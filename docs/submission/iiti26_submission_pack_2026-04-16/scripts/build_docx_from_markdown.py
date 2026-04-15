#!/usr/bin/env python3
"""Lightweight Markdown -> DOCX converter for article drafts."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Mm, Pt


HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
IMAGE_RE = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")
BULLET_RE = re.compile(r"^\s*[-*]\s+(.*)$")
NUMBER_RE = re.compile(r"^\s*\d+\.\s+(.*)$")


def apply_run_font(paragraph, font_name: str, font_size_pt: int) -> None:
    for run in paragraph.runs:
        run.font.name = font_name
        run.font.size = Pt(font_size_pt)


def clean_inline_md(text: str) -> str:
    text = text.replace("`", "")
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    return text.strip()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    block = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        block.append(lines[i].strip())
        i += 1

    rows = []
    for idx, row in enumerate(block):
        cols = [c.strip() for c in row.strip("|").split("|")]
        if idx == 1 and all(set(c) <= {"-", ":"} for c in cols):
            continue
        rows.append(cols)
    return rows, i


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(ncols):
            val = row[c_idx] if c_idx < len(row) else ""
            cell = table.cell(r_idx, c_idx)
            cell.text = clean_inline_md(val)
    doc.add_paragraph("")


def convert(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25.4)
    section.bottom_margin = Mm(25.4)
    section.left_margin = Mm(25.4)
    section.right_margin = Mm(25.4)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    i = 0
    in_code = False
    code_lines: list[str] = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph("\n".join(code_lines))
                apply_run_font(p, "Courier New", 9)
                doc.add_paragraph("")
                in_code = False
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        m = HEADING_RE.match(stripped)
        if m:
            level = len(m.group(1))
            text = clean_inline_md(m.group(2))
            if level == 1:
                p = doc.add_heading(text, level=0)
                apply_run_font(p, "Times New Roman", 14)
            else:
                p = doc.add_heading(text, level=min(level, 4))
                apply_run_font(p, "Times New Roman", 12 if level == 2 else 11)
            i += 1
            continue

        m = IMAGE_RE.match(stripped)
        if m:
            caption = clean_inline_md(m.group(1))
            rel_path = m.group(2).strip()
            img_path = (md_path.parent / rel_path).resolve()
            if img_path.exists():
                doc.add_picture(str(img_path), width=Inches(6.2))
                if caption:
                    cp = doc.add_paragraph(caption)
                    apply_run_font(cp, "Times New Roman", 10)
            else:
                p = doc.add_paragraph(f"[Missing image: {rel_path}]")
                apply_run_font(p, "Times New Roman", 10)
            doc.add_paragraph("")
            i += 1
            continue

        if stripped.startswith("|"):
            rows, i2 = parse_table(lines, i)
            add_table(doc, rows)
            i = i2
            continue

        bm = BULLET_RE.match(line)
        if bm:
            p = doc.add_paragraph(clean_inline_md(bm.group(1)), style="List Bullet")
            apply_run_font(p, "Times New Roman", 11)
            i += 1
            continue

        nm = NUMBER_RE.match(line)
        if nm:
            p = doc.add_paragraph(clean_inline_md(nm.group(1)), style="List Number")
            apply_run_font(p, "Times New Roman", 11)
            i += 1
            continue

        paragraph_lines = [stripped]
        j = i + 1
        while j < len(lines):
            nxt = lines[j].strip()
            if (
                not nxt
                or nxt == "---"
                or nxt.startswith("|")
                or nxt.startswith("```")
                or HEADING_RE.match(nxt)
                or IMAGE_RE.match(nxt)
                or BULLET_RE.match(lines[j])
                or NUMBER_RE.match(lines[j])
            ):
                break
            paragraph_lines.append(nxt)
            j += 1
        text = clean_inline_md(" ".join(paragraph_lines))
        p = doc.add_paragraph(text)
        apply_run_font(p, "Times New Roman", 11)
        i = j

    doc.save(docx_path)


def main() -> None:
    if len(sys.argv) != 3:
        print("Usage: build_docx_from_markdown.py <input.md> <output.docx>")
        raise SystemExit(1)
    md_path = Path(sys.argv[1]).resolve()
    docx_path = Path(sys.argv[2]).resolve()
    convert(md_path, docx_path)
    print(f"Saved: {docx_path}")


if __name__ == "__main__":
    main()
